"""Small reusable deterministic and LLM-backed tool implementations."""

from __future__ import annotations

import csv
import json
import math
import operator
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any


def _records(arguments: dict[str, Any]) -> list[dict[str, Any]]:
    records = arguments.get("records", [])
    if not isinstance(records, list) or any(not isinstance(row, dict) for row in records):
        raise ValueError("records must be a list of objects")
    return records


def tabular_file_reader(arguments: dict[str, Any]) -> dict[str, Any]:
    path = Path(arguments["path"])
    if not path.is_file():
        raise FileNotFoundError(f"Tabular file not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
    elif suffix == ".json":
        rows = json.loads(path.read_text(encoding="utf-8"))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[arguments.get("sheet_name", workbook.sheetnames[0])]
        values = list(worksheet.values)
        headers = [str(value) for value in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:]]
        workbook.close()
    else:
        raise ValueError("Supported tabular formats are CSV, JSON, and XLSX")
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError("Tabular data must contain object rows")
    return {"columns": list(rows[0]) if rows else [], "rows": rows, "source": str(path)}


def file_exporter(arguments: dict[str, Any]) -> dict[str, Any]:
    rows = arguments.get("rows", [])
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError("rows must be a list of objects")
    path = Path(arguments["path"])
    path.parent.mkdir(parents=True, exist_ok=True)
    fmt = arguments.get("format", path.suffix.lstrip(".").lower() or "json")
    if fmt == "json":
        path.write_text(json.dumps(rows, indent=2, default=str), encoding="utf-8")
    elif fmt == "csv":
        columns = list(arguments.get("columns") or (list(rows[0]) if rows else []))
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)
    else:
        raise ValueError("Supported export formats are CSV and JSON")
    return {"path": str(path), "rows_exported": len(rows), "format": fmt}


def schema_normalizer(arguments: dict[str, Any]) -> dict[str, Any]:
    columns = arguments.get("columns", [])
    mapping = {
        str(column): re.sub(r"[^a-z0-9]+", "_", str(column).strip().lower()).strip("_")
        for column in columns
    }
    return {"mapping": mapping, "columns": list(mapping.values())}


def record_validator(arguments: dict[str, Any]) -> dict[str, Any]:
    rows = _records(arguments)
    required = arguments.get("required_fields", [])
    valid, invalid = [], []
    for row in rows:
        missing = [field for field in required if row.get(field) in (None, "")]
        (invalid if missing else valid).append(
            {"record": row, "missing_fields": missing} if missing else row
        )
    return {"valid": valid, "invalid": invalid, "valid_count": len(valid), "invalid_count": len(invalid)}


def record_lookup(arguments: dict[str, Any]) -> dict[str, Any]:
    rows = _records(arguments)
    criteria = arguments.get("criteria", {})
    matches = [row for row in rows if all(row.get(key) == value for key, value in criteria.items())]
    return {"matches": matches, "found": bool(matches)}


def shipment_lookup(arguments: dict[str, Any]) -> dict[str, Any]:
    rows = _records(arguments)
    order_id = arguments.get("order_id")
    matches = [row for row in rows if row.get("order_id") == order_id]
    return {"matches": matches, "found": bool(matches)}


def record_matcher(arguments: dict[str, Any]) -> dict[str, Any]:
    left, right = _records({"records": arguments.get("left", [])}), _records({"records": arguments.get("right", [])})
    key = arguments["key"]
    index = {row.get(key): row for row in right}
    matched, unmatched = [], []
    for row in left:
        other = index.get(row.get(key))
        (matched if other is not None else unmatched).append({"left": row, "right": other} if other else row)
    return {"matched": matched, "unmatched": unmatched}


def text_normalizer(arguments: dict[str, Any]) -> dict[str, Any]:
    values = arguments.get("values", [])
    normalized = [re.sub(r"\s+", " ", str(value).strip().lower()) for value in values]
    return {"values": normalized}


def similarity_scorer(arguments: dict[str, Any]) -> dict[str, Any]:
    left, right = str(arguments.get("left", "")), str(arguments.get("right", ""))
    return {"score": SequenceMatcher(None, left.casefold(), right.casefold()).ratio()}


def numeric_calculator(arguments: dict[str, Any]) -> dict[str, Any]:
    operation = arguments["operation"]
    if operation == "success_failure_rate":
        rows = _records(arguments)
        total = len(rows)
        failures = sum(1 for row in rows if str(row.get("status", "")).lower() == "failed")
        return {
            "total": total,
            "failures": failures,
            "successes": total - failures,
            "failure_rate": failures / total * 100 if total else 0,
            "operation": operation,
        }
    if operation == "average_field":
        rows = _records(arguments)
        field = arguments["field"]
        values = [float(row[field]) for row in rows]
        return {
            "field": field,
            "average": sum(values) / len(values) if values else 0,
            "count": len(values),
            "operation": operation,
        }
    if operation == "reorder_quantity":
        rows = _records(arguments)
        results = [
            {**row, "reorder_quantity": max(float(row["minimum_stock"]) - float(row["current_stock"]), 0)}
            for row in rows
            if "minimum_stock" in row and "current_stock" in row
        ]
        return {"result": results, "operation": operation}
    if operation == "percentage_difference" and arguments.get("records"):
        results = []
        for match in _records(arguments):
            left, right = match["left"], match["right"]
            internal, vendor = float(left["internal_price"]), float(right["vendor_price"])
            results.append({
                **left,
                "vendor_price": vendor,
                "percentage_difference": abs(vendor - internal) / abs(internal) * 100,
            })
        return {"result": results, "operation": operation}
    values = [float(value) for value in arguments.get("values", [])]
    if operation == "sum":
        result = sum(values)
    elif operation == "average":
        result = sum(values) / len(values) if values else 0
    elif operation == "percentage_difference":
        if len(values) != 2 or values[0] == 0:
            raise ValueError("percentage_difference requires two values and a non-zero first value")
        result = abs(values[1] - values[0]) / abs(values[0]) * 100
    elif operation == "subtract":
        result = values[0] - values[1]
    else:
        raise ValueError(f"Unsupported numeric operation: {operation}")
    return {"result": result, "operation": operation}


def threshold_checker(arguments: dict[str, Any]) -> dict[str, Any]:
    operations = {"<": operator.lt, "<=": operator.le, ">": operator.gt, ">=": operator.ge, "==": operator.eq}
    comparison = arguments.get("operator", ">")
    if comparison not in operations:
        raise ValueError(f"Unsupported threshold operator: {comparison}")
    if "records" in arguments:
        value_field = arguments["value_field"]
        threshold_field = arguments.get("threshold_field")
        threshold_value = arguments.get("threshold")
        flagged = [
            row for row in _records(arguments)
            if operations[comparison](
                float(row[value_field]),
                float(row[threshold_field]) if threshold_field else float(threshold_value),
            )
        ]
        return {"flagged": bool(flagged), "records": flagged, "operator": comparison}
    value, threshold = arguments["value"], arguments["threshold"]
    return {"value": value, "threshold": threshold, "flagged": operations[comparison](value, threshold)}


def priority_scorer(arguments: dict[str, Any]) -> dict[str, Any]:
    record = arguments["record"]
    weights = arguments.get("weights", {})
    score = sum(float(record.get(field, 0)) * float(weight) for field, weight in weights.items())
    return {"score": score, "record": record}


def capacity_checker(arguments: dict[str, Any]) -> dict[str, Any]:
    workload = float(arguments["workload"])
    capacity = float(arguments["capacity"])
    return {"available": workload < capacity, "remaining": capacity - workload}


def candidate_ranker(arguments: dict[str, Any]) -> dict[str, Any]:
    candidates = _records({"records": arguments.get("candidates", [])})
    field = arguments.get("score_field", "score")
    ranked = sorted(candidates, key=lambda row: float(row.get(field, 0)), reverse=True)
    return {"ranked": ranked}


def text_validator(arguments: dict[str, Any]) -> dict[str, Any]:
    text = str(arguments.get("text", ""))
    required = arguments.get("required_terms", [])
    missing = [term for term in required if term not in text]
    return {"valid": not missing, "missing_terms": missing}


def error_analyzer(arguments: dict[str, Any]) -> dict[str, Any]:
    errors = [str(row.get("error", "")) for row in _records(arguments) if row.get("error")]
    counts: dict[str, int] = {}
    for error in errors:
        counts[error] = counts.get(error, 0) + 1
    return {"error_counts": counts, "total_errors": len(errors)}


def report_builder(arguments: dict[str, Any]) -> dict[str, Any]:
    sections = arguments.get("sections", {})
    if not isinstance(sections, dict):
        raise ValueError("sections must be an object")
    return {"title": arguments.get("title", "Report"), "sections": sections}


def _llm_json(llm_client: Any, instruction: str, facts: Any) -> dict[str, Any]:
    if llm_client is None:
        raise ValueError("An LLM client is required for this tool")
    response = llm_client.complete(
        f"{instruction}\nUse only these facts. Do not invent missing information.\n"
        f"Facts: {json.dumps(facts, default=str)}\nReturn valid JSON only."
    )
    result = json.loads(response)
    if not isinstance(result, dict):
        raise ValueError("LLM response must be a JSON object")
    return result


def build_default_tools(*, llm_client: Any = None) -> dict[str, Any]:
    """Return the complete generic tool set for registration."""
    tools = {
        "tabular_file_reader": tabular_file_reader,
        "file_exporter": file_exporter,
        "schema_normalizer": schema_normalizer,
        "record_validator": record_validator,
        "record_lookup": record_lookup,
        "shipment_lookup": shipment_lookup,
        "record_matcher": record_matcher,
        "text_normalizer": text_normalizer,
        "similarity_scorer": similarity_scorer,
        "numeric_calculator": numeric_calculator,
        "threshold_checker": threshold_checker,
        "priority_scorer": priority_scorer,
        "capacity_checker": capacity_checker,
        "candidate_ranker": candidate_ranker,
        "text_validator": text_validator,
        "error_analyzer": error_analyzer,
        "report_builder": report_builder,
    }
    tools.update({
        "text_classifier": lambda args: _llm_json(llm_client, "Classify the supplied text.", args),
        "text_generator": lambda args: _llm_json(llm_client, "Generate the requested text fields.", args),
        "text_summarizer": lambda args: _llm_json(llm_client, "Summarize the supplied facts.", args),
        "recommendation_generator": lambda args: _llm_json(llm_client, "Generate recommendations from the supplied facts.", args),
    })
    return tools
