"""Load workflow definitions from Excel workbooks."""

from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.models import Workflow


class WorkflowLoaderError(ValueError):
    """Raised when a workflow workbook cannot be loaded or validated."""


def _is_empty_row(values: Iterable[Any]) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in values)


def load_workflows(
    file_path: str | Path = Path("data") / "workflows.xlsx",
    required_columns: tuple[str, ...] = ("Workflow_ID", "Workflow_Name"),
) -> list[Workflow]:
    """Load non-empty workflow rows from the first worksheet in an Excel file."""
    path = Path(file_path)
    if not path.is_file():
        raise WorkflowLoaderError(f"Excel workflow file not found: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, InvalidFileException, ValueError, TypeError) as exc:
        raise WorkflowLoaderError(f"Malformed Excel workflow file '{path}': {exc}") from exc

    try:
        if not workbook.worksheets:
            raise WorkflowLoaderError(f"Excel workflow file '{path}' contains no worksheets")

        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration as exc:
            raise WorkflowLoaderError(f"Excel workflow file '{path}' is empty") from exc

        if _is_empty_row(header_values):
            raise WorkflowLoaderError(f"Excel workflow file '{path}' has no header row")

        headers = [
            value.strip() if isinstance(value, str) else value
            for value in header_values
        ]
        if any(not isinstance(header, str) or not header for header in headers):
            raise WorkflowLoaderError(
                f"Excel workflow file '{path}' has a blank or invalid header"
            )
        if len(set(headers)) != len(headers):
            raise WorkflowLoaderError(f"Excel workflow file '{path}' has duplicate headers")

        missing_columns = [column for column in required_columns if column not in headers]
        if missing_columns:
            missing = ", ".join(missing_columns)
            raise WorkflowLoaderError(
                f"Excel workflow file '{path}' is missing required columns: {missing}"
            )

        workflows: list[Workflow] = []
        for row_values in rows:
            values = list(row_values)
            if _is_empty_row(values):
                continue
            values.extend([None] * (len(headers) - len(values)))
            fields = dict(zip(headers, values[: len(headers)]))
            workflows.append(
                Workflow(
                    workflow_id=fields["Workflow_ID"],
                    workflow_name=fields["Workflow_Name"],
                    fields=fields,
                )
            )
        return workflows
    finally:
        workbook.close()
